"""
File Tools — small desktop utility (tkinter, no extra GUI dependency).

Tabs:
  1. Bulk rename   : find/replace, prefix/suffix, numbering, with live preview
  2. Merge PDFs    : pick several PDFs, reorder, merge into one
  3. CSV -> Excel  : convert one or many CSV files to formatted .xlsx

Run:  python file_tools.py
"""

from __future__ import annotations

import csv
import os
import re
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfWriter


# =========================================================================== #
# Tab 1 — Bulk rename
# =========================================================================== #
class RenameTab(ttk.Frame):
    def __init__(self, master: tk.Misc) -> None:
        super().__init__(master, padding=10)
        self.folder = tk.StringVar()
        self.find = tk.StringVar()
        self.replace = tk.StringVar()
        self.prefix = tk.StringVar()
        self.suffix = tk.StringVar()
        self.numbering = tk.BooleanVar(value=False)
        self.start = tk.IntVar(value=1)
        self.ext_filter = tk.StringVar(value="*")

        row = 0
        ttk.Label(self, text="Folder").grid(row=row, column=0, sticky="w")
        ttk.Entry(self, textvariable=self.folder, width=50).grid(row=row, column=1, sticky="we")
        ttk.Button(self, text="Browse…", command=self.pick_folder).grid(row=row, column=2, padx=4)

        for label, var in (("Find", self.find), ("Replace with", self.replace),
                           ("Prefix", self.prefix), ("Suffix", self.suffix),
                           ("Extension filter (e.g. jpg)", self.ext_filter)):
            row += 1
            ttk.Label(self, text=label).grid(row=row, column=0, sticky="w")
            ttk.Entry(self, textvariable=var).grid(row=row, column=1, sticky="we")

        row += 1
        ttk.Checkbutton(self, text="Add numbering  (file_001, file_002 …)", variable=self.numbering).grid(row=row, column=1, sticky="w")
        ttk.Spinbox(self, from_=0, to=99999, textvariable=self.start, width=6).grid(row=row, column=2)

        row += 1
        self.tree = ttk.Treeview(self, columns=("old", "new"), show="headings", height=12)
        self.tree.heading("old", text="Current name")
        self.tree.heading("new", text="New name")
        self.tree.grid(row=row, column=0, columnspan=3, sticky="nsew", pady=8)
        self.rowconfigure(row, weight=1)
        self.columnconfigure(1, weight=1)

        row += 1
        ttk.Button(self, text="Preview", command=self.preview).grid(row=row, column=1, sticky="e")
        ttk.Button(self, text="Rename files", command=self.apply).grid(row=row, column=2)

        for var in (self.find, self.replace, self.prefix, self.suffix, self.numbering, self.start, self.ext_filter):
            var.trace_add("write", lambda *_: self.preview())

    def pick_folder(self) -> None:
        if d := filedialog.askdirectory():
            self.folder.set(d)
            self.preview()

    def plan(self) -> list[tuple[Path, str]]:
        folder = Path(self.folder.get())
        if not folder.is_dir():
            return []
        ext = self.ext_filter.get().strip().lstrip(".").lower()
        files = sorted(p for p in folder.iterdir() if p.is_file() and (ext in ("", "*") or p.suffix.lower().lstrip(".") == ext))
        width = max(3, len(str(len(files) + self.start.get())))
        result = []
        for i, path in enumerate(files):
            stem = path.stem
            if self.find.get():
                stem = stem.replace(self.find.get(), self.replace.get())
            stem = f"{self.prefix.get()}{stem}{self.suffix.get()}"
            if self.numbering.get():
                stem = f"{stem}_{i + self.start.get():0{width}d}"
            result.append((path, stem + path.suffix))
        return result

    def preview(self) -> None:
        self.tree.delete(*self.tree.get_children())
        for path, new in self.plan():
            self.tree.insert("", "end", values=(path.name, new))

    def apply(self) -> None:
        plan = [(p, n) for p, n in self.plan() if p.name != n]
        if not plan:
            messagebox.showinfo("Nothing to do", "No file name would change.")
            return
        targets = [n for _, n in plan]
        if len(set(targets)) != len(targets):
            messagebox.showerror("Conflict", "Two files would get the same name. Enable numbering.")
            return
        if not messagebox.askyesno("Confirm", f"Rename {len(plan)} file(s)?"):
            return
        # two-pass rename to avoid collisions with existing names
        tmp = [(p, p.with_name(f"__tmp_{i}_{p.name}")) for i, (p, _) in enumerate(plan)]
        for p, t in tmp:
            p.rename(t)
        for (p, t), (_, new) in zip(tmp, plan):
            t.rename(p.with_name(new))
        self.preview()
        messagebox.showinfo("Done", f"{len(plan)} file(s) renamed.")


# =========================================================================== #
# Tab 2 — Merge PDFs
# =========================================================================== #
class MergeTab(ttk.Frame):
    def __init__(self, master: tk.Misc) -> None:
        super().__init__(master, padding=10)
        self.listbox = tk.Listbox(self, selectmode="single", height=14)
        self.listbox.grid(row=0, column=0, columnspan=4, sticky="nsew", pady=(0, 8))
        self.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)

        ttk.Button(self, text="Add PDFs…", command=self.add).grid(row=1, column=0, sticky="w")
        ttk.Button(self, text="↑", width=3, command=lambda: self.move(-1)).grid(row=1, column=1)
        ttk.Button(self, text="↓", width=3, command=lambda: self.move(1)).grid(row=1, column=2)
        ttk.Button(self, text="Remove", command=self.remove).grid(row=1, column=3)
        ttk.Button(self, text="Merge → Save as…", command=self.merge).grid(row=2, column=0, columnspan=4, pady=8, sticky="we")

    def add(self) -> None:
        for f in filedialog.askopenfilenames(filetypes=[("PDF", "*.pdf")]):
            self.listbox.insert("end", f)

    def remove(self) -> None:
        if sel := self.listbox.curselection():
            self.listbox.delete(sel[0])

    def move(self, delta: int) -> None:
        sel = self.listbox.curselection()
        if not sel:
            return
        i, j = sel[0], sel[0] + delta
        if 0 <= j < self.listbox.size():
            item = self.listbox.get(i)
            self.listbox.delete(i)
            self.listbox.insert(j, item)
            self.listbox.selection_set(j)

    def merge(self) -> None:
        files = self.listbox.get(0, "end")
        if len(files) < 2:
            messagebox.showwarning("Need more files", "Add at least two PDF files.")
            return
        out = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF", "*.pdf")], initialfile="merged.pdf")
        if not out:
            return
        writer = PdfWriter()
        for f in files:
            writer.append(f)
        with open(out, "wb") as fh:
            writer.write(fh)
        messagebox.showinfo("Done", f"{len(files)} PDFs merged into:\n{out}")


# =========================================================================== #
# Tab 3 — CSV -> Excel
# =========================================================================== #
def sniff_dialect(sample: str) -> csv.Dialect | type[csv.Dialect]:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def coerce(value: str):
    """Turn numeric-looking strings into numbers so Excel can compute on them."""
    v = value.strip()
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    if re.fullmatch(r"-?\d+[.,]\d+", v):
        return float(v.replace(",", "."))
    return value


def csv_to_xlsx(src: Path, dst: Path) -> int:
    with open(src, newline="", encoding="utf-8-sig") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        rows = list(csv.reader(fh, dialect=sniff_dialect(sample)))
    wb = Workbook()
    ws = wb.active
    ws.title = src.stem[:31]
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value if r == 1 else coerce(value))
    if rows:
        for c in range(1, len(rows[0]) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 50)
    wb.save(dst)
    return len(rows)


class CsvTab(ttk.Frame):
    def __init__(self, master: tk.Misc) -> None:
        super().__init__(master, padding=10)
        ttk.Label(self, text="Convert one or many CSV files into formatted Excel workbooks\n"
                             "(auto-detects ; , tab | separators, bold header, filters, column widths).",
                  justify="left").grid(row=0, column=0, sticky="w", pady=(0, 10))
        ttk.Button(self, text="Choose CSV files and convert…", command=self.convert).grid(row=1, column=0, sticky="we")
        self.log = tk.Text(self, height=14, state="disabled")
        self.log.grid(row=2, column=0, sticky="nsew", pady=8)
        self.rowconfigure(2, weight=1)
        self.columnconfigure(0, weight=1)

    def write(self, line: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", line + "\n")
        self.log.configure(state="disabled")
        self.log.see("end")

    def convert(self) -> None:
        files = filedialog.askopenfilenames(filetypes=[("CSV", "*.csv *.tsv *.txt")])
        for f in files:
            src = Path(f)
            dst = src.with_suffix(".xlsx")
            try:
                n = csv_to_xlsx(src, dst)
                self.write(f"OK  {src.name} -> {dst.name} ({n} rows)")
            except Exception as exc:  # show the error, keep going with the next file
                self.write(f"ERR {src.name}: {exc}")


# =========================================================================== #
def main() -> None:
    root = tk.Tk()
    root.title("File Tools")
    root.geometry("720x520")
    try:
        ttk.Style().theme_use("vista" if os.name == "nt" else "clam")
    except tk.TclError:
        pass
    nb = ttk.Notebook(root)
    nb.add(RenameTab(nb), text="  Bulk rename  ")
    nb.add(MergeTab(nb), text="  Merge PDFs  ")
    nb.add(CsvTab(nb), text="  CSV → Excel  ")
    nb.pack(fill="both", expand=True)
    root.mainloop()


if __name__ == "__main__":
    main()
